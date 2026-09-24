#!/usr/bin/env python3
import argparse
import csv
import hashlib
import json
import re
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

FIELDS = [
    "fiscalYear", "month", "endUser", "positionDesignation", "fundCode", "fundClass", "prNumber", "modality", "item", "quantity", "unitOfIssue", "unitBudget", "estimatedTotal", "unitLcrb", "total", "purpose", "office", "lonReceivedDate", "bacAwardApprovedDate", "poContractDate", "poContractDateReceived", "deliveryDate", "supplier", "iarDate", "sectionCodeDepartment", "releasedDate", "status", "remarks", "obrNumber", "rfqsPrinted", "salesChargeInvoice"
]

MONTH_RE = re.compile(r"^(\d+)\.\s*SVP\s+2025-([A-Za-z]+)")

def clean(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).replace("\u00a0", " ").strip()
    return text or None

def number(value):
    value = clean(value)
    if value is None:
        return None
    try:
        return float(value.replace(",", ""))
    except ValueError:
        return value

def choose_monthly_files(extracted):
    grouped = defaultdict(list)
    for path in extracted.rglob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        match = MONTH_RE.match(path.name)
        if match:
            grouped[(int(match.group(1)), match.group(2).lower())].append(path)
    selected = []
    duplicates = []
    for key, paths in sorted(grouped.items()):
        preferred = sorted(paths, key=lambda p: ("(1)" not in p.stem, len(p.name), p.name))[0]
        selected.append(preferred)
        duplicates.extend(p for p in paths if p != preferred)
    return selected, duplicates

def row_to_record(row, workbook_name, source_row, occurrence):
    values = [clean(v) for v in row]
    values += [None] * (31 - len(values))
    if str(values[0]) not in {"2025", "2025.0"} or not values[6]:
        return None
    record = dict(zip(FIELDS, values[:31]))
    record["fiscalYear"] = 2025
    for key in ("quantity", "unitBudget", "estimatedTotal", "unitLcrb", "total"):
        record[key] = number(record[key])
    record["sourceWorkbook"] = workbook_name
    record["sourceRow"] = source_row
    identity = "|".join(str(record.get(key) or "").strip().lower() for key in ("fiscalYear", "month", "prNumber", "item", "quantity", "unitOfIssue", "unitBudget", "estimatedTotal", "unitLcrb", "total", "supplier"))
    record["recordKey"] = hashlib.sha256(f"{identity}|{occurrence}".encode()).hexdigest()[:64]
    return record

def normalize(zip_path, output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory() as temp:
        temp_path = Path(temp)
        with zipfile.ZipFile(zip_path) as archive:
            archive.extractall(temp_path)
        selected, duplicates = choose_monthly_files(temp_path)
        records = []
        for path in selected:
            wb = load_workbook(path, read_only=True, data_only=True)
            if "DATA" not in wb.sheetnames:
                continue
            ws = wb["DATA"]
            header_found = False
            occurrences = Counter()
            for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
                values = list(row)
                if not header_found:
                    if any(str(value).strip().lower() == "pr#" for value in values if value is not None):
                        header_found = True
                    continue
                if not any(value not in (None, "") for value in values):
                    continue
                provisional = [clean(value) for value in values]
                if str(provisional[0]) not in {"2025", "2025.0"} or not provisional[6]:
                    continue
                identity = "|".join(str(provisional[index] if index < len(provisional) else "").strip().lower() for index in (0, 1, 6, 8, 9, 10, 11, 12, 13, 14, 22))
                occurrences[identity] += 1
                record = row_to_record(values, path.name, row_number, occurrences[identity])
                if record:
                    records.append(record)
        records.sort(key=lambda item: (item["prNumber"], item["sourceRow"], item["item"] or ""))
        json_path = output_dir / "pmr-2025-normalized.json"
        csv_path = output_dir / "pmr-2025-normalized.csv"
        manifest_path = output_dir / "pmr-2025-import-manifest.json"
        json_path.write_text(json.dumps(records, indent=2, ensure_ascii=False), encoding="utf-8")
        with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDS + ["recordKey", "sourceWorkbook", "sourceRow"])
            writer.writeheader()
            writer.writerows(records)
        manifest = {
            "sourceArchive": Path(zip_path).name,
            "selectedWorkbooks": [p.name for p in selected],
            "ignoredDuplicateWorkbooks": [p.name for p in duplicates],
            "sourceSheet": "DATA",
            "includedFiscalYear": 2025,
            "excludedYears": [2021, 2022, 2024],
            "recordCount": len(records),
            "purchaseRequestCount": len({r["prNumber"] for r in records}),
            "months": sorted({r["month"] for r in records if r["month"]}),
        }
        manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
        print(json.dumps(manifest, indent=2, ensure_ascii=False))

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("archive")
    parser.add_argument("output_dir")
    args = parser.parse_args()
    normalize(Path(args.archive), Path(args.output_dir))
